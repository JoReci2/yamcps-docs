"""PowerPoint text extraction and atomic ingestion with the native index."""

import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import Mock, patch

from pptx import Presentation
from pptx.exc import PackageNotFoundError
from pptx.util import Inches
from yamcps.rag import Index

from yamcps_docs.ingest import _extract, ingest


class PowerPointTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.path = self.root / "evidence.pptx"
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_textbox(0, 0, Inches(4), Inches(1)).text = "REQ-42\nSauvegarde"
        table = slide.shapes.add_table(2, 2, 0, Inches(1), Inches(4), Inches(2)).table
        table.cell(0, 0).merge(table.cell(0, 1))
        table.cell(0, 0).text = "Exigences"
        table.cell(1, 0).text = "Statut"
        table.cell(1, 1).text = "Approuvé"
        group = slide.shapes.add_group_shape()
        nested = group.shapes.add_group_shape()
        nested.shapes.add_textbox(0, 0, Inches(1), Inches(1)).text = "Groupe imbriqué"
        slide.shapes.add_connector(1, 0, 0, 100, 100)
        slide.notes_slide.notes_text_frame.text = "Note non indexée"
        blank = presentation.slides.add_slide(presentation.slide_layouts[6])
        blank.shapes.add_table(1, 2, 0, 0, Inches(2), Inches(1))
        last = presentation.slides.add_slide(presentation.slide_layouts[6])
        last.shapes.add_textbox(0, 0, Inches(1), Inches(1)).text = "Conclusion"
        presentation.save(self.path)

    def test_text_tables_nested_groups_and_slide_numbers(self):
        text = _extract(self.path)
        self.assertEqual(
            text,
            "Slide: 1\nREQ-42\nSauvegarde\nExigences\nStatut | Approuvé\n"
            "Groupe imbriqué\nSlide: 3\nConclusion",
        )
        upper = self.path.with_suffix(".PPTX")
        self.path.rename(upper)
        self.assertEqual(_extract(upper), text)

    def test_ingestion_citations_and_invalid_snapshot_rollback(self):
        index = Index(self.root / "index.sqlite3")
        manifest = self.root / "manifest.json"
        manifest.write_text(
            json.dumps(
                {
                    "corpus": "slides",
                    "sources": [{"path": self.path.name, "document_id": "review", "revision": "1"}],
                }
            )
        )
        self.assertEqual(ingest(index, manifest)["updated"], 1)
        self.assertEqual(ingest(index, manifest)["unchanged"], 1)
        hit = index.search("REQ-42", "slides")[0]
        text = index.get_document("slides", "review", "1").text
        self.assertEqual(text[hit.start : hit.end], hit.text)
        self.assertEqual(index.get_passage(hit.passage_id).text, hit.text)
        self.path.write_bytes(b"invalid")
        with self.assertRaises(PackageNotFoundError):
            ingest(index, manifest)
        self.assertEqual(index.get_document("slides", "review", "1").text, text)
        empty = Presentation()
        empty.slides.add_slide(empty.slide_layouts[6])
        empty.save(self.path)
        with self.assertRaisesRegex(ValueError, "empty"):
            ingest(index, manifest)
        self.assertEqual(index.get_document("slides", "review", "1").text, text)

    def test_extracted_character_limit(self):
        shape = Mock(spec=["has_text_frame", "text_frame"])
        shape.has_text_frame = True
        shape.text_frame.text = "x" * 4_000_001
        module = Mock()
        module.Presentation.return_value.slides = [Mock(shapes=[shape])]
        with patch("yamcps_docs.ingest.importlib.import_module", return_value=module):
            with self.assertRaisesRegex(ValueError, "4 million"):
                _extract(self.path)
