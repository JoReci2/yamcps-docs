"""Excel extraction, evidence retrieval and snapshot rollback."""

import json
import tempfile
import unittest
from datetime import datetime
from pathlib import Path
from unittest.mock import Mock, patch

from openpyxl import Workbook
from yamcps.rag import Index

from yamcps_docs.ingest import _excel_text, _extract, ingest


class ExcelTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.path = self.root / "evidence.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Coûts"
        sheet.append(["Référence", "Montant", "Actif", "Date"])
        sheet.append(["DEV-42", 123.5, False, datetime(2026, 10, 24)])
        sheet["AA4"] = "=SUM(B2:B3)"
        workbook.create_sheet("Vide")
        hidden = workbook.create_sheet("Masquée")
        hidden.sheet_state = "hidden"
        hidden["C3"] = "preuve cachée"
        workbook.save(self.path)
        workbook.close()

    def test_xlsx_preserves_sheets_coordinates_types_and_formulas(self):
        text = _extract(self.path)
        self.assertIn("Sheet: Coûts", text)
        self.assertIn("A2: DEV-42 | B2: 123.5 | C2: False | D2: 2026-10-24T00:00:00", text)
        self.assertIn("AA4: =SUM(B2:B3)", text)
        self.assertIn("Sheet: Masquée\nC3: preuve cachée", text)
        self.assertNotIn("Sheet: Vide", text)
        upper = self.path.with_suffix(".XLSX")
        self.path.rename(upper)
        self.assertEqual(text, _extract(upper))

    def test_xls_saved_values_dates_booleans_and_errors(self):
        text = _extract(Path(__file__).parent / "fixtures" / "evidence.xls")
        self.assertIn("Sheet: Coûts", text)
        self.assertIn("A2: DEV-42", text)
        self.assertIn("B2: 123.5", text)
        self.assertIn("C2: False", text)
        self.assertIn("D2: 2026-10-24T00:00:00", text)
        self.assertIn("E2: #DIV/0!", text)
        self.assertIn("Sheet: Suite\nC3: preuve", text)

    def test_excel_ingestion_search_citations_and_rollback(self):
        index = Index(self.root / "index.sqlite3")
        manifest = self.root / "manifest.json"
        source = {"path": self.path.name, "document_id": "budget", "revision": "1"}
        manifest.write_text(json.dumps({"corpus": "finance", "sources": [source]}))
        self.assertEqual(ingest(index, manifest)["updated"], 1)
        self.assertEqual(ingest(index, manifest)["unchanged"], 1)
        hit = index.search("DEV-42", "finance")[0]
        text = index.get_document("finance", "budget", "1").text
        self.assertEqual(text[hit.start : hit.end], hit.text)
        self.assertIn("Sheet: Coûts", text)
        for content in (b"corrupt workbook", None):
            if content is None:
                empty = Workbook()
                empty.save(self.path)
                empty.close()
            else:
                self.path.write_bytes(content)
            with self.assertRaises(Exception):
                ingest(index, manifest)
            self.assertEqual(index.get_document("finance", "budget", "1").text, text)

    def test_limits_and_blank_cells(self):
        self.assertEqual(_excel_text([("empty", [[None, ""]])]), "")
        self.assertEqual(_excel_text([("zero", [[0, False]])]), "Sheet: zero\nA1: 0 | B1: False")
        with self.assertRaisesRegex(ValueError, "4 million"):
            _excel_text([("large", [["x" * 4_000_001]])])
        with self.assertRaisesRegex(ValueError, "1 million"):
            _excel_text([("sparse", ([None] for _ in range(1_000_001)))])

    def test_xlsx_closes_on_extraction_failure(self):
        workbook = Mock()
        workbook.worksheets = [Mock(title="large")]
        workbook.worksheets[0].iter_rows.return_value = [["x" * 4_000_001]]
        module = Mock()
        module.load_workbook.return_value = workbook
        with patch("yamcps_docs.ingest.importlib.import_module", return_value=module):
            with self.assertRaises(ValueError):
                _extract(self.path)
        workbook.close.assert_called_once()
