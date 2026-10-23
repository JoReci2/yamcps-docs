"""Offline ingestion and real MCP clients for the docs connector."""

import asyncio
import contextlib
import io
import json
import os
import socket
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import Mock, patch

from mcp import Client, StdioServerParameters
from yamcps.rag import Document, Index
from yamcps.server import create_server

from yamcps_docs import DocsConnector
from yamcps_docs.cli import main
from yamcps_docs.ingest import _extract, ingest


class SiblingConnector:
    id = "sibling"

    def register(self, server):
        @server.tool()
        def sibling_echo(message: str) -> str:
            return message


def docs_and_sibling():
    with patch("yamcps.server.load", return_value=[DocsConnector(), SiblingConnector()]):
        return create_server()


class IngestTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.index = Index(self.root / "index.sqlite3")
        self.manifest = self.root / "manifest.json"
        (self.root / "SSS.md").write_text("SSS-REQ-0042 requires backup.", encoding="utf-8")
        self.source = {"path": "SSS.md", "document_id": "SSS", "revision": "1"}
        self.write_manifest([self.source])

    def write_manifest(self, sources):
        self.manifest.write_text(json.dumps({"corpus": "project/specs", "sources": sources}))

    def test_ingest_snapshot_and_cli(self):
        self.assertEqual(ingest(self.index, self.manifest)["updated"], 1)
        self.assertEqual(ingest(self.index, self.manifest)["unchanged"], 1)
        with patch.dict(os.environ, {"YAMCPS_DOCS_INDEX": str(self.index.path)}, clear=True):
            with (
                patch("sys.argv", ["yamcps-docs", "ingest", str(self.manifest)]),
                contextlib.redirect_stdout(io.StringIO()) as output,
            ):
                main()
            self.assertEqual(json.loads(output.getvalue())["unchanged"], 1)
        self.write_manifest([])
        self.assertEqual(ingest(self.index, self.manifest)["deleted"], 1)

    def test_invalid_sources_leave_previous_snapshot_intact(self):
        ingest(self.index, self.manifest)
        (self.root / "link.md").symlink_to(self.root / "SSS.md")
        (self.root / "empty.txt").write_text("")
        (self.root / "unsupported.xyz").write_text("x")
        for name in ["../outside.md", "/etc/passwd", "link.md", "empty.txt", "unsupported.xyz"]:
            with self.subTest(name=name):
                self.write_manifest([{**self.source, "path": name}])
                with self.assertRaises(ValueError):
                    ingest(self.index, self.manifest)
                self.assertEqual(self.index.status("project/specs")["documents"], 1)
        self.write_manifest([{**self.source, "path": "missing.md"}])
        with self.assertRaises(FileNotFoundError):
            ingest(self.index, self.manifest)
        with (
            patch("sys.argv", ["yamcps-docs", "ingest", str(self.manifest)]),
            contextlib.redirect_stderr(io.StringIO()) as error,
        ):
            with self.assertRaises(SystemExit) as caught:
                main()
        self.assertEqual(caught.exception.code, 1)
        self.assertIn("Ingestion failed", error.getvalue())
        self.assertNotIn(str(self.root), error.getvalue())

    def test_source_size_and_directory_limits(self):
        self.write_manifest([{**self.source, "path": "."}])
        with self.assertRaises(ValueError):
            ingest(self.index, self.manifest)
        self.write_manifest([self.source])
        with patch("yamcps_docs.ingest._extract", return_value="x" * 4_000_001):
            with self.assertRaises(ValueError):
                ingest(self.index, self.manifest)
        huge = self.root / "huge.md"
        with huge.open("wb") as stream:
            stream.truncate(21 * 1024 * 1024)
        self.write_manifest([{**self.source, "path": "huge.md"}])
        with self.assertRaisesRegex(ValueError, "20 MiB"):
            ingest(self.index, self.manifest)

    def test_optional_pdf_docx_extractors(self):
        page = Mock()
        page.extract_text.return_value = "PDF text"
        pdf = Mock()
        pdf.PdfReader.return_value.pages = [page]
        with patch("yamcps_docs.ingest.importlib.import_module", return_value=pdf):
            self.assertEqual(_extract(self.root / "test.pdf"), "PDF text")
        word = Mock()
        word.Document.return_value.paragraphs = [Mock(text="Word text")]
        word.Document.return_value.tables = [Mock(rows=[Mock(cells=[Mock(text="cell")])])]
        with patch("yamcps_docs.ingest.importlib.import_module", return_value=word):
            self.assertEqual(_extract(self.root / "test.docx"), "Word text\ncell")

    def test_real_pdf_docx_formats(self):
        from docx import Document
        from pypdf import PdfWriter

        document = Document()
        document.add_paragraph("Word requirement")
        document.add_table(rows=1, cols=1).cell(0, 0).text = "table value"
        document.save(str(self.root / "test.docx"))
        self.assertIn("table value", _extract(self.root / "test.docx"))
        pdf = PdfWriter()
        pdf.add_blank_page(width=100, height=100)
        pdf.write(str(self.root / "blank.pdf"))
        self.assertEqual(_extract(self.root / "blank.pdf"), "")


class DocsIntegrationTests(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        root = Path(self.temp.name)
        self.path = root / "index.sqlite3"
        (root / "SSS.md").write_text("SSS-REQ-0042 requires backup.")
        manifest = root / "manifest.json"
        manifest.write_text(
            json.dumps(
                {
                    "corpus": "project/specs",
                    "sources": [{"path": "SSS.md", "document_id": "SSS", "revision": "1"}],
                }
            )
        )
        ingest(Index(self.path), manifest)
        environment = patch.dict(os.environ, {"YAMCPS_DOCS_INDEX": str(self.path)}, clear=False)
        environment.start()
        self.addCleanup(environment.stop)

    async def exercise(self, target):
        async with Client(target) as client:
            tools = await client.list_tools()
            self.assertTrue(
                {
                    "docs_search",
                    "docs_get_document",
                    "docs_get_passage",
                    "docs_index_status",
                }.issubset({t.name for t in tools.tools})
            )
            result = await client.call_tool(
                "docs_search", {"query": "SSS-REQ-0042", "corpus": "project/specs"}
            )
            hits = result.structured_content["result"]
            hit = hits[0]
            self.assertEqual(hit["revision"], "1")
            content = await client.read_resource(hit["source_uri"])
            self.assertIn("requires backup", content.contents[0].text)
            passage = await client.read_resource("docs://passage/" + hit["passage_id"])
            self.assertIn("SSS-REQ-0042", passage.contents[0].text)
            fetched = await client.call_tool("docs_get_passage", {"passage_id": hit["passage_id"]})
            self.assertEqual(fetched.structured_content["text"], hit["text"])
            document = await client.call_tool(
                "docs_get_document",
                {"corpus": "project/specs", "document_id": "SSS", "revision": "1"},
            )
            self.assertEqual(document.structured_content["source_uri"], hit["source_uri"])
            status = await client.call_tool("docs_index_status", {"corpus": "project/specs"})
            self.assertEqual(status.structured_content["documents"], 1)
            invalid = await client.call_tool("docs_get_passage", {"passage_id": "../missing"})
            self.assertTrue(invalid.is_error)
            self.assertNotIn(str(self.path), str(invalid.content))
            skill = await client.read_resource("skill://docs-research/SKILL.md")
            self.assertIn("docs_search", skill.contents[0].text)

    async def test_memory_and_coexistence(self):
        await self.exercise(docs_and_sibling())

    async def test_stdio(self):
        await self.exercise(
            StdioServerParameters(
                command=sys.executable,
                args=["-m", "yamcps", "serve", "--connector", "docs"],
                env={"YAMCPS_DOCS_INDEX": str(self.path)},
            )
        )

    async def test_http(self):
        with socket.socket() as sock:
            sock.bind(("127.0.0.1", 0))
            port = sock.getsockname()[1]
        with tempfile.TemporaryFile() as output:
            process = await asyncio.create_subprocess_exec(
                sys.executable,
                "-m",
                "yamcps",
                "serve",
                "--connector",
                "docs",
                "--transport",
                "http",
                "--port",
                str(port),
                stdout=output,
                stderr=output,
            )
            try:
                async with asyncio.timeout(15):
                    while True:
                        if process.returncode is not None:
                            self.fail("HTTP docs process exited")
                        try:
                            _, writer = await asyncio.open_connection("127.0.0.1", port)
                            writer.close()
                            await writer.wait_closed()
                            break
                        except OSError:
                            await asyncio.sleep(0.05)
                await self.exercise(f"http://127.0.0.1:{port}/mcp")
            finally:
                if process.returncode is None:
                    process.terminate()
                await asyncio.wait_for(process.wait(), timeout=10)

    async def test_escaped_resource_identity_and_readonly_annotations(self):
        document = Document(
            corpus="space %2F/é", document_id="a/b%20c", revision="1 #2", text="Evidence"
        )
        Index(self.path).sync(document.corpus, [document])
        async with Client(create_server(["docs"])) as client:
            content = await client.read_resource(document.source_uri)
            self.assertEqual(content.contents[0].text, "Evidence")
            tools = (await client.list_tools()).tools
            for tool in tools:
                self.assertTrue(tool.annotations.read_only_hint)
                self.assertFalse(tool.annotations.destructive_hint)

    async def test_missing_index_does_not_break_other_plugins(self):
        with patch.dict(os.environ, {"YAMCPS_DOCS_INDEX": str(self.path.parent / "missing")}):
            server = docs_and_sibling()
            async with Client(server) as client:
                result = await client.call_tool("sibling_echo", {"message": "ok"})
                self.assertEqual(result.structured_content["result"], "ok")
                result = await client.call_tool("docs_index_status", {"corpus": "x"})
                self.assertTrue(result.is_error)
